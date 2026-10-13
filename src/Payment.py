import os
import re
import pandas as pd
from openpyxl import load_workbook

# ---------------- CONFIG ----------------
INPUT_FOLDER  = r"C:/path/to/input_folder"
OUTPUT_FILE   = r"C:/path/to/output.xlsx"     # same workbook that has Payment, Cost, ... sheets
SHEET_NAME    = "Payment"

# Input column numbers (1-based, as in your description)
NAME_COL      = 2
REF_COL       = 3
DIRECT_COL    = 5
AMOUNT_COL_1  = 12
AMOUNT_COL_2  = 13

MARKER             = "remittance report"
STOP_AMOUNT        = 109676.29
STOP_SECOND_VALUE  = 0.00

OUT_COLUMNS = ["DEFENDANT'S NAME", "REF NO.", "AMOUNT PAID",
               "DIRECT PAYMENT", "CLIENT SHARE", "BW, LLP SHARE", "Notes"]
# ----------------------------------------


def to_num(v):
    """'1,09,676.29' / '$200.00' / 200 -> float. Returns None if not numeric."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[,$\s]", "", str(v))
    try:
        return float(s)
    except ValueError:
        return None


def is_blank(v):
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == ""


def is_stop_row(row):
    """True when two consecutive cells are 109676.29 and 0.00."""
    nums = [to_num(c) for c in row]
    for a, b in zip(nums, nums[1:]):
        if a is not None and b is not None \
           and abs(a - STOP_AMOUNT) < 0.005 and abs(b - STOP_SECOND_VALUE) < 0.005:
            return True
    return False


def process_file(path):
    """Read one input file into a list of output rows (the 'datatable')."""
    df = pd.read_excel(path, header=None, dtype=object)
    rows_out, started = [], False

    for _, r in df.iterrows():
        row = list(r.values)

        # 1) don't take anything until the "Remittance Report" row has passed
        if not started:
            if any(not is_blank(c) and MARKER in str(c).lower() for c in row):
                started = True          # this row itself is NOT taken
            continue

        # 2) stop at the total row (109676.29 / 0.00)
        if is_stop_row(row):
            break

        # 3) skip junk rows (blank, "Current Balance: $1,906.95", etc.)
        name = row[NAME_COL - 1] if len(row) >= NAME_COL else None
        ref  = row[REF_COL - 1]  if len(row) >= REF_COL else None
        if is_blank(name) or is_blank(ref):
            continue
        if str(name).strip().lower().startswith("current balance"):
            continue

        a1 = to_num(row[AMOUNT_COL_1 - 1]) if len(row) >= AMOUNT_COL_1 else None
        a2 = to_num(row[AMOUNT_COL_2 - 1]) if len(row) >= AMOUNT_COL_2 else None
        direct = to_num(row[DIRECT_COL - 1]) if len(row) >= DIRECT_COL else None

        rows_out.append([
            str(name).strip(),
            str(ref).strip(),
            round((a1 or 0) + (a2 or 0), 2),   # AMOUNT PAID = col 12 + col 13
            direct if direct is not None else 0.0,
            None,                               # CLIENT SHARE  (empty for now)
            None,                               # BW, LLP SHARE (empty for now)
            None,                               # Notes
        ])
    return rows_out


def build_datatable():
    all_rows = []
    for f in sorted(os.listdir(INPUT_FOLDER)):
        if f.startswith("~$") or not f.lower().endswith((".xlsx", ".xlsm", ".xls")):
            continue
        rows = process_file(os.path.join(INPUT_FOLDER, f))
        print(f"{f}: {len(rows)} rows")
        all_rows.extend(rows)
    return pd.DataFrame(all_rows, columns=OUT_COLUMNS)   # <- the DataTable


def append_to_excel(dt):
    wb = load_workbook(OUTPUT_FILE)          # other sheets are preserved untouched
    ws = wb[SHEET_NAME]

    # locate header row / first column from the template
    hdr_row = hdr_col = total_row = None
    for row in ws.iter_rows():
        for c in row:
            v = str(c.value).strip().upper() if c.value is not None else ""
            if v == "DEFENDANT'S NAME" and hdr_row is None:
                hdr_row, hdr_col = c.row, c.column
            if v == "TOTAL" and total_row is None:
                total_row = c.row
    if hdr_row is None:
        raise ValueError("Header row not found in Payment sheet")

    # find the first empty row after the header
    r = hdr_row + 1
    while any(not is_blank(ws.cell(r, hdr_col + i).value) for i in range(len(OUT_COLUMNS))):
        r += 1
    first_data = hdr_row + 1

    # write all rows in one pass
    for rec in dt.itertuples(index=False):
        for i, val in enumerate(rec):
            ws.cell(r, hdr_col + i, None if pd.isna(val) else val)
        r += 1
    last_data = r - 1

    # refresh the Total row (AMOUNT PAID, DIRECT, CLIENT SHARE, BW SHARE)
    if total_row:
        for i in range(2, 6):
            col = hdr_col + i
            L = ws.cell(1, col).column_letter
            ws.cell(total_row, col, f"=SUM({L}{first_data}:{L}{last_data})")

    wb.save(OUTPUT_FILE)


if __name__ == "__main__":
    dt = build_datatable()
    print(dt.head(), f"\nTotal rows: {len(dt)}")
    append_to_excel(dt)
    print("Done.")